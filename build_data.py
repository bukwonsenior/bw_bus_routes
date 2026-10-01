# -*- coding: utf-8 -*-
"""
셔틀버스_노선관리.xlsx  ->  route_data.js / routes.json / stops.csv / routes_수정본.csv

담당자는 엑셀(또는 홈페이지 편집기)만 고친다. GitHub Actions가 이 스크립트를 돌린다.
- '정류장' 시트: 번호 | 정류장명 | 정류장번호 | 위도 | 경도 | 주소 | 비고 | 좌표확인
    정류장번호(5자리 BIS)만 넣으면 위도·경도는 '원주시_버스정류장_좌표.csv'에서 자동으로 채운다(엑셀에도 써 넣음).
- '코스' 시트: 게시 | 코스 | 종류(코스정보/정차) | 정류장명 | 시각 | 설명 | 색상
검사: 시간 역행·형식 오류·없는 정류장 → 멈춤(홈페이지 그대로). 속도 이상 → 경고만.
"""
import csv, json, math, os, re, sys, hashlib
import openpyxl

OUT  = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(OUT, "셔틀버스_노선관리.xlsx")
BIS_CSV = os.path.join(OUT, "원주시_버스정류장_좌표.csv")
KNOWN_ID = {"오전 1코스": "am1", "오전 2코스": "am2", "오후 코스": "pm"}   # 기존 코스 id 유지
TIME = re.compile(r"^([01]\d|2[0-3]):[0-5]\d$")
COLOR = re.compile(r"^#[0-9A-Fa-f]{6}$")
DEFAULT_COLORS = ["#E8442E", "#1F6FEB", "#1D8348", "#8E44AD", "#D35400", "#16A085"]

errors, warns = [], []
def s(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()
def as_time(v):
    if hasattr(v, "hour") and hasattr(v, "minute"): return f"{v.hour:02d}:{v.minute:02d}"
    if isinstance(v, float) and 0 <= v < 1:   # 엑셀이 시각을 숫자로 저장한 경우
        m = round(v * 1440); return f"{m // 60:02d}:{m % 60:02d}"
    t = s(v)
    m = re.match(r"^(\d{1,2}):(\d{2})$", t)
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else t
def mins(t): h, m = t.split(":"); return int(h) * 60 + int(m)
def hav(a, b, c, e):
    p = math.radians; dlat = p(c - a); dlon = p(e - b)
    x = math.sin(dlat / 2) ** 2 + math.cos(p(a)) * math.cos(p(c)) * math.sin(dlon / 2) ** 2
    return 2 * 6371.0 * math.asin(math.sqrt(x))

def head_ix(ws, need, sheet):
    head = [s(c.value) for c in ws[4]]
    for h in need:
        if h not in head: errors.append(f"'{sheet}' 시트 4행에 '{h}' 칸이 없습니다. 칸 이름을 바꾸지 마세요.")
    return {h: head.index(h) for h in head if h}

def main():
    if not os.path.exists(XLSX):
        print(f"::error::{os.path.basename(XLSX)} 파일이 없습니다."); sys.exit(1)
    wb = openpyxl.load_workbook(XLSX)
    for n in ("정류장", "코스"):
        if n not in wb.sheetnames: print(f"::error::'{n}' 시트가 없습니다."); sys.exit(1)

    bis = {}
    if os.path.exists(BIS_CSV):
        for r in csv.DictReader(open(BIS_CSV, encoding="utf-8-sig")):
            if r.get("모바일단축번호") and r["모바일단축번호"] != "0":
                bis[r["모바일단축번호"]] = (float(r["위도"]), float(r["경도"]))

    # ---- 정류장 ----
    ws = wb["정류장"]
    ix = head_ix(ws, ["번호", "정류장명", "정류장번호", "위도", "경도", "주소", "비고"], "정류장")
    if errors: [print("::error::" + e) for e in errors]; sys.exit(1)
    col = lambda row, h: s(row[ix[h]].value) if h in ix and ix[h] < len(row) else ""
    stops, byname, used_ids, pending_id = [], {}, set(), []
    for r in range(5, ws.max_row + 1):
        row = ws[r]
        if not any(s(c.value) for c in row): continue
        name = col(row, "정류장명")
        if not name: errors.append(f"정류장 {r}행: 정류장명이 비어 있습니다"); continue
        if name in byname: errors.append(f"정류장 {r}행: '{name}' 이름이 {byname[name]['_row']}행에도 있습니다 (이름은 겹치면 안 돼요)"); continue
        no, b = col(row, "정류장번호").replace(".0", ""), None
        lat, lng = col(row, "위도"), col(row, "경도")
        if not (lat and lng):
            if no and no in bis:
                lat, lng = bis[no]
                ws.cell(r, ix["위도"] + 1, lat); ws.cell(r, ix["경도"] + 1, lng)
                b = True
            else:
                errors.append(f"정류장 {r}행: '{name}' 위치를 모릅니다 → 정류장번호(5자리)를 넣거나 위도·경도를 직접 넣어 주세요"
                              + (f" (정류장번호 {no} 는 원주시 목록에 없음)" if no else "")); continue
        try: lat, lng = float(lat), float(lng)
        except ValueError: errors.append(f"정류장 {r}행: 위도·경도는 숫자로 (지금: {lat}, {lng})"); continue
        if not (37.20 <= lat <= 37.55 and 127.75 <= lng <= 128.15):
            warns.append(f"정류장 {r}행: '{name}' 좌표가 원주시 범위 밖입니다 ({lat}, {lng})")
        sid = col(row, "번호")
        st = {"_row": r, "id": int(float(sid)) if sid.replace(".", "").isdigit() else None, "name": name, "lat": lat, "lng": lng,
              "bis": no, "addr": col(row, "주소"), "conf": col(row, "좌표확인") or "confirmed", "note": col(row, "비고")}
        if b: st["_filled"] = True
        if st["id"] is not None:
            if st["id"] in used_ids: errors.append(f"정류장 {r}행: 번호 {st['id']} 이(가) 겹칩니다")
            used_ids.add(st["id"])
        else: pending_id.append(st)
        stops.append(st); byname[name] = st
    nxt = max(used_ids or {0}) + 1
    for st in pending_id:   # 번호 빈 새 정류장 → 자동 번호 (엑셀에도 기록)
        st["id"] = nxt; ws.cell(st["_row"], ix["번호"] + 1, nxt); nxt += 1; st["_filled"] = True

    # ---- 코스 ----
    wc = wb["코스"]
    cx = head_ix(wc, ["게시", "코스", "종류", "정류장명", "시각", "설명", "색상"], "코스")
    if errors: [print("::error::" + e) for e in errors]; sys.exit(1)
    routes, rmap = [], {}
    for r in range(5, wc.max_row + 1):
        row = wc[r]
        if not any(s(c.value) for c in row): continue
        g = lambda h: s(row[cx[h]].value) if cx[h] < len(row) else ""
        if g("게시").upper() in ("X", "×"): continue
        cname, kind = g("코스"), g("종류")
        if not cname: errors.append(f"코스 {r}행: 코스 칸이 비어 있습니다"); continue
        if cname not in rmap:
            rmap[cname] = {"id": KNOWN_ID.get(cname, f"c{len(routes) + 1}"), "name": cname, "sub": "",
                           "color": DEFAULT_COLORS[len(routes) % len(DEFAULT_COLORS)], "seq": []}
            routes.append(rmap[cname])
        R = rmap[cname]
        if kind == "코스정보":
            R["sub"] = g("설명")
            c = g("색상")
            if c:
                if COLOR.match(c): R["color"] = c.upper()
                else: errors.append(f"코스 {r}행: 색상은 #E8442E 처럼 써 주세요 (지금: {c})")
        elif kind == "정차":
            nm, t = g("정류장명"), as_time(row[cx["시각"]].value)
            if nm not in byname: errors.append(f"코스 {r}행: 정류장 '{nm}' 이(가) 정류장 시트에 없습니다"); continue
            if not TIME.match(t): errors.append(f"코스 {r}행: 시각은 07:50 처럼 써 주세요 (지금: '{t}')"); continue
            if R["seq"] and mins(t) < mins(R["seq"][-1][1]):
                errors.append(f"코스 {r}행: {cname} '{nm}' {t} 이(가) 앞 정류장({R['seq'][-1][0]['name']} {R['seq'][-1][1]})보다 빠릅니다 (시간 역행)")
            R["seq"].append((byname[nm], t))
        else:
            errors.append(f"코스 {r}행: 종류는 코스정보 / 정차 중 하나 (지금: '{kind}')")
    for R in routes:
        if len(R["seq"]) < 2: errors.append(f"코스 '{R['name']}' 에 정차 줄이 2개 이상 있어야 합니다")
        for (a, ta), (b, tb) in zip(R["seq"], R["seq"][1:]):   # 속도 점검(경고)
            km, dt = hav(a["lat"], a["lng"], b["lat"], b["lng"]), mins(tb) - mins(ta)
            if dt > 0 and km / (dt / 60) > 75: warns.append(f"{R['name']} {a['name']}→{b['name']}: {km:.1f}km를 {dt}분 (속도 과다 — 시각 확인)")
            if dt == 0 and km > 0.6: warns.append(f"{R['name']} {a['name']}→{b['name']}: 같은 시각인데 {km:.1f}km 떨어져 있음")

    if errors:
        print("\n엑셀에서 고쳐야 할 곳이 있습니다. 노선도는 그대로 둡니다.\n")
        for e in errors: print("  ✗ " + e)
        print(f"::error::엑셀 오류 {len(errors)}건"); sys.exit(1)

    if any(st.get("_filled") for st in stops):
        wb.save(XLSX); print("  · 엑셀에 좌표·번호 자동 채움:", ", ".join(st["name"] for st in stops if st.get("_filled")))

    # ---- 출력 (기존과 같은 형식) ----
    with open(os.path.join(OUT, "stops.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["stop_id", "정류장명", "위도", "경도", "정류장번호", "주소", "좌표신뢰도", "비고"])
        for st in stops: w.writerow([st["id"], st["name"], st["lat"], st["lng"], st["bis"], st["addr"], st["conf"], st["note"]])
    data = {"stops": [{"id": st["id"], "name": st["name"], "lat": st["lat"], "lng": st["lng"], "bis": st["bis"],
                       "addr": st["addr"], "conf": st["conf"], "note": st["note"]} for st in stops], "routes": []}
    for R in routes:
        data["routes"].append({"id": R["id"], "name": R["name"], "sub": R["sub"], "color": R["color"],
            "stops": [{"seq": i + 1, "id": st["id"], "name": st["name"], "time": t, "lat": st["lat"], "lng": st["lng"], "bis": st["bis"]}
                      for i, (st, t) in enumerate(R["seq"])]})
    js = ("// 노선 데이터 - 자동 생성 파일입니다. 셔틀버스_노선관리.xlsx(또는 홈페이지 편집기)를 고치세요.\n"
          "const ROUTE_DATA = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n")
    open(os.path.join(OUT, "route_data.js"), "w", encoding="utf-8").write(js)
    json.dump(data, open(os.path.join(OUT, "routes.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    with open(os.path.join(OUT, "routes_수정본.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["코스", "순번", "정류장명", "시간", "stop_id", "정류장번호", "좌표신뢰도"])
        for R in routes:
            for i, (st, t) in enumerate(R["seq"]): w.writerow([R["name"], i + 1, st["name"], t, st["id"], st["bis"], st["conf"]])
    ver = hashlib.md5(js.encode("utf-8")).hexdigest()[:8]   # 캐시 무효화
    for fn in ("index.html", "print.html"):
        p = os.path.join(OUT, fn)
        if not os.path.exists(p): continue
        h = open(p, encoding="utf-8").read()
        n = re.sub(r"route_data\.js(\?v=[0-9a-f]+)?", "route_data.js?v=" + ver, h)
        if n != h: open(p, "w", encoding="utf-8").write(n)
    for w_ in warns: print("::warning::" + w_)
    print(f"✓ 정류장 {len(stops)}곳 | " + " / ".join(f"{R['name']} {len(R['seq'])}개소" for R in routes) + f" | 캐시 {ver}")

if __name__ == "__main__":
    main()
